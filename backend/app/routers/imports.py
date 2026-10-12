"""Bulk product/stock import: upload -> validate & preview -> commit (background, resumable, audited)."""
import asyncio
import csv
import io
import logging
import re
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation

from bson import ObjectId
from fastapi import APIRouter, Depends, File, Request, UploadFile

from ..audit import write_audit
from ..db import db, run_atomic
from ..security import require_owner
from ..services import cache
from ..services.stock import apply_delta, build_txn, price_fields, sync_inventory_snapshot
from ..utils import AppError, gen_code, jsonable, norm, not_found, now, oid, to_dec, tokens

log = logging.getLogger("jaykisan.imports")
router = APIRouter(prefix="/imports", tags=["imports"])
MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024
COLUMNS = ["sku", "name", "category", "brand", "unit", "barcode", "supplier", "purchase_price", "selling_price", "min_stock",
           "batch_no", "expiry_date", "description", "warehouse_code", "quantity"]


def _read_rows(filename: str, raw: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        it = wb.active.iter_rows(values_only=True)
        header = [norm(str(h or "")).replace(" ", "_") for h in next(it, [])]
        return [{header[i]: ("" if v is None else str(v)).strip() for i, v in enumerate(r) if i < len(header)} for r in it if any(v not in (None, "") for v in r)]
    if name.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        reader.fieldnames = [norm(h).replace(" ", "_") for h in reader.fieldnames or []]
        return [{k: (v or "").strip() for k, v in r.items() if k} for r in reader if any((v or "").strip() for v in r.values())]
    raise AppError(422, "UNSUPPORTED_FILE", "Upload a .csv or .xlsx file")


def _price(v, field, errors):
    if v == "":
        return None
    try:
        d = Decimal(v.replace(",", ""))
        if d < 0:
            raise InvalidOperation
        return str(d.quantize(Decimal("0.01")))
    except (InvalidOperation, ValueError):
        errors.append(f"{field} must be a positive number")
        return None


def _int(v, field, errors, positive=False):
    if v == "":
        return None
    try:
        n = int(float(v))
        if n < 0 or (positive and n == 0) or n > 100_000_000:
            raise ValueError
        return n
    except ValueError:
        errors.append(f"{field} must be a {'positive ' if positive else ''}whole number")
        return None


@router.post("/products", summary="Upload CSV/XLSX and get a validated preview (nothing is saved yet)")
async def upload(request: Request, file: UploadFile = File(...), owner: dict = Depends(require_owner)):
    raw = await file.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise AppError(413, "FILE_TOO_LARGE", "File is larger than 5 MB. Split it into smaller files.")
    rows = _read_rows(file.filename or "", raw)
    if not rows:
        raise AppError(422, "EMPTY_FILE", "The file has no data rows")
    if len(rows) > MAX_ROWS:
        raise AppError(422, "TOO_MANY_ROWS", f"Maximum {MAX_ROWS} rows per import. Split the file.")
    if "sku" not in rows[0] or "name" not in rows[0]:
        raise AppError(422, "MISSING_COLUMNS", "The file must include 'sku' and 'name' columns")
    warehouses = {w["code"]: w for w in await db.warehouses.find({"active": True}).to_list(1000)}
    skus = list({r.get("sku", "").upper() for r in rows if r.get("sku")})
    existing = {p["sku_norm"]: p for p in await db.products.find({"sku_norm": {"$in": skus}}, {"sku_norm": 1, "barcode": 1}).to_list(len(skus) + 1)}
    barcodes = list({r.get("barcode") for r in rows if r.get("barcode")})
    barcode_owner = {p["barcode"]: p["sku_norm"] for p in await db.products.find({"barcode": {"$in": barcodes}}, {"barcode": 1, "sku_norm": 1}).to_list(len(barcodes) + 1)}
    job_id, t = ObjectId(), now()
    seen_sku, seen_bc, docs = set(), set(), []
    stats = {"total": len(rows), "valid": 0, "invalid": 0, "creates": 0, "updates": 0, "stock_rows": 0}
    for i, r in enumerate(rows, start=2):
        errors: list = []
        sku = r.get("sku", "").strip()
        sku_n = sku.upper()
        if not sku or len(sku) > 64:
            errors.append("sku is required (max 64 chars)")
        elif sku_n in seen_sku:
            errors.append("duplicate sku in file")
        if not r.get("name") and sku_n not in existing:
            errors.append("name is required for new products")
        bc = r.get("barcode", "")
        if bc:
            if bc in seen_bc:
                errors.append("duplicate barcode in file")
            elif barcode_owner.get(bc) not in (None, sku_n):
                errors.append("barcode already used by another product")
        exp = r.get("expiry_date", "")
        if exp:
            try:
                date.fromisoformat(exp[:10])
            except ValueError:
                errors.append("expiry_date must be YYYY-MM-DD")
        pp = _price(r.get("purchase_price", ""), "purchase_price", errors)
        sp = _price(r.get("selling_price", ""), "selling_price", errors)
        ms = _int(r.get("min_stock", ""), "min_stock", errors)
        qty = _int(r.get("quantity", ""), "quantity", errors, positive=True)
        wcode = r.get("warehouse_code", "").upper()
        if qty and wcode not in warehouses:
            errors.append(f"warehouse_code '{wcode}' not found" if wcode else "warehouse_code is required when quantity is given")
        seen_sku.add(sku_n)
        if bc:
            seen_bc.add(bc)
        action = "update" if sku_n in existing else "create"
        valid = not errors
        stats["valid" if valid else "invalid"] += 1
        if valid:
            stats["creates" if action == "create" else "updates"] += 1
            stats["stock_rows"] += 1 if qty else 0
        docs.append({"job_id": job_id, "row_no": i, "valid": valid, "errors": errors, "action": action, "done": False,
                     "data": {**{c: r.get(c, "") for c in COLUMNS}, "sku": sku, "purchase_price": pp, "selling_price": sp,
                              "min_stock": ms, "quantity": qty, "warehouse_code": wcode},
                     "expires_at": t + timedelta(days=2)})
    await db.import_rows.insert_many(docs)
    job = {"_id": job_id, "job_no": gen_code("IM"), "filename": file.filename, "status": "preview", "stats": stats, "user_id": owner["_id"],
           "created_at": t, "expires_at": t + timedelta(days=2), "result": None}
    await db.import_jobs.insert_one(job)
    return await _job_view(job)


async def _job_view(job: dict) -> dict:
    sample = await db.import_rows.find({"job_id": job["_id"]}).sort([("valid", 1), ("row_no", 1)]).limit(100).to_list(100)
    return jsonable({"id": job["_id"], "job_no": job["job_no"], "filename": job.get("filename"), "status": job["status"], "stats": job["stats"],
                     "result": job.get("result"), "rows": [{"row_no": s["row_no"], "valid": s["valid"], "errors": s["errors"], "action": s["action"],
                                                             "sku": s["data"]["sku"], "name": s["data"].get("name"), "quantity": s["data"].get("quantity"),
                                                             "warehouse_code": s["data"].get("warehouse_code")} for s in sample]})


@router.get("/{job_id}")
async def get_job(job_id: str, _owner: dict = Depends(require_owner)):
    job = await db.import_jobs.find_one({"_id": oid(job_id)})
    if not job:
        raise not_found("Import job")
    return await _job_view(job)


async def _category_path(path: str, cache_map: dict):
    parts = [p.strip() for p in re.split(r">|/", path) if p.strip()]
    parent, ancestors = None, []
    for name in parts:
        key = (parent, norm(name))
        if key not in cache_map:
            doc = await db.categories.find_one({"parent_id": parent, "name_lc": norm(name)})
            if not doc:
                t = now()
                doc = {"name": name, "name_lc": norm(name), "parent_id": parent, "ancestors": list(ancestors), "depth": len(ancestors), "active": True, "created_at": t, "updated_at": t}
                doc["_id"] = (await db.categories.insert_one(doc)).inserted_id
            cache_map[key] = doc["_id"]
        ancestors.append(cache_map[key])
        parent = cache_map[key]
    return (parent, ancestors) if parts else (None, [])


async def _process(job: dict, owner: dict):
    warehouses = {w["code"]: w for w in await db.warehouses.find({"active": True}).to_list(1000)}
    cats: dict = {}
    ok = failed = stock_ok = 0
    async for row in db.import_rows.find({"job_id": job["_id"], "valid": True, "done": False}).sort("row_no", 1):
        d = row["data"]
        try:
            fields = {k: d[k] for k in ("brand", "unit", "supplier", "batch_no", "description") if d.get(k)}
            if d.get("name"):
                fields.update({"name": d["name"], "name_lc": norm(d["name"])})
            for k in ("purchase_price", "selling_price"):
                if d.get(k) is not None:
                    fields[k] = to_dec(d[k])
            if d.get("min_stock") is not None:
                fields["min_stock"] = d["min_stock"]
            if d.get("barcode"):
                fields["barcode"] = d["barcode"]
            if d.get("expiry_date"):
                fields["expiry_date"] = d["expiry_date"][:10]
            if d.get("category"):
                fields["category_id"], ids = await _category_path(d["category"], cats)
                fields["category_ids"] = ids
            sku_n, t = d["sku"].upper(), now()
            current = await db.products.find_one({"sku_norm": sku_n})
            merged = {**(current or {}), **fields}
            fields["search_tokens"] = tokens(merged.get("name"), merged.get("brand"), d["sku"])
            fields["updated_at"] = t
            if current:
                product = await db.products.find_one_and_update({"_id": current["_id"]}, {"$set": fields}, return_document=True)
                await sync_inventory_snapshot(product)
            else:
                product = {"sku": d["sku"], "sku_norm": sku_n, "unit": "pcs", "min_stock": 0, "purchase_price": None, "selling_price": None,
                           "category_id": None, "category_ids": [], "active": True, "created_at": t, "created_by": owner["_id"], **fields}
                product["_id"] = (await db.products.insert_one(product)).inserted_id
            if d.get("quantity"):
                wh, qty = warehouses[d["warehouse_code"]], d["quantity"]

                async def fn(s, product=product, wh=wh, qty=qty):
                    before, after, _inv = await apply_delta(s, product, wh["_id"], qty)
                    pf = price_fields(qty, None, product.get("purchase_price"))
                    txn = build_txn(type_="STOCK_IN", product=product, warehouse=wh, user=owner, change=qty, before=before, after=after,
                                    note=f"Bulk import {job['job_no']}", import_job_id=job["_id"], unit_cost=pf["unit_cost"], total_cost=pf["total_cost"])
                    await db.stock_transactions.insert_one(txn, session=s)
                    await db.import_rows.update_one({"_id": row["_id"]}, {"$set": {"done": True}}, session=s)

                await run_atomic(fn)
                stock_ok += 1
            else:
                await db.import_rows.update_one({"_id": row["_id"]}, {"$set": {"done": True}})
            ok += 1
        except Exception as e:  # noqa: BLE001
            failed += 1
            msg = getattr(e, "message", None) or ("duplicate barcode/sku" if "E11000" in str(e) else "failed")
            await db.import_rows.update_one({"_id": row["_id"]}, {"$set": {"errors": [msg]}})
            log.warning("import row %s failed: %s", row["row_no"], e)
    result = {"imported": ok, "failed": failed, "stock_rows": stock_ok}
    await db.import_jobs.update_one({"_id": job["_id"]}, {"$set": {"status": "completed", "result": result, "completed_at": now()}})
    cache.clear()
    await write_audit(owner, "BULK_IMPORT", "import_job", job["_id"], after=result, meta={"job_no": job["job_no"], "filename": job.get("filename")})


@router.post("/{job_id}/commit", summary="Apply the validated rows (runs in background; poll GET /imports/{id})")
async def commit(job_id: str, owner: dict = Depends(require_owner)):
    job = await db.import_jobs.find_one_and_update(
        {"_id": oid(job_id), "status": "preview"}, {"$set": {"status": "committing", "started_at": now()}}, return_document=True
    )
    if not job:
        raise AppError(409, "IMPORT_NOT_PENDING", "This import was already committed or has expired")
    asyncio.create_task(_process(job, owner))
    return await _job_view(job)
